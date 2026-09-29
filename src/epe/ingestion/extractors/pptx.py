"""PPTX extractor using python-pptx, with graceful fallback."""

from __future__ import annotations

from pathlib import Path

from epe.ingestion.extractors import ExtractedDocument


def extract(path: Path) -> ExtractedDocument:
    warnings: list[str] = []
    try:
        from pptx import Presentation  # type: ignore[import-not-found]
    except ImportError as e:  # pragma: no cover
        warnings.append(f"python-pptx not installed: {e}. Falling back to raw-text read.")
        return ExtractedDocument(
            path=path,
            text=path.read_text(encoding="utf-8", errors="replace"),
            extractor="pptx-fallback",
            warnings=warnings,
        )

    prs = Presentation(path)
    pages: list[str] = []
    for i, slide in enumerate(prs.slides, start=1):
        bits: list[str] = [f"--- Slide {i} ---"]
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs)
                    if line.strip():
                        bits.append(line)
        pages.append("\n".join(bits))

    return ExtractedDocument(
        path=path,
        text="\n\n".join(pages),
        pages=pages,
        structure={"slide_count": len(prs.slides)},
        extractor="pptx",
        warnings=warnings,
    )
