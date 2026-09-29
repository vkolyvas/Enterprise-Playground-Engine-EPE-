"""XLSX extractor using openpyxl, with graceful fallback."""

from __future__ import annotations

from pathlib import Path

from epe.ingestion.extractors import ExtractedDocument


def extract(path: Path) -> ExtractedDocument:
    warnings: list[str] = []
    try:
        import openpyxl  # type: ignore[import-not-found]
    except ImportError as e:  # pragma: no cover
        warnings.append(f"openpyxl not installed: {e}. Falling back to raw-text read.")
        return ExtractedDocument(
            path=path,
            text=path.read_text(encoding="utf-8", errors="replace"),
            extractor="xlsx-fallback",
            warnings=warnings,
        )

    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    tables: list[list[list[str]]] = []
    parts: list[str] = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f"# Sheet: {sheet_name}")
        sheet_rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            sheet_rows.append(["" if v is None else str(v) for v in row])
        if sheet_rows:
            tables.append(sheet_rows)
            parts.append("\n".join(" | ".join(r) for r in sheet_rows))

    return ExtractedDocument(
        path=path,
        text="\n\n".join(parts),
        structure={"sheet_count": len(wb.sheetnames), "sheet_names": list(wb.sheetnames)},
        tables=tables,
        extractor="xlsx",
        warnings=warnings,
    )
