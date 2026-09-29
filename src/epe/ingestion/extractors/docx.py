"""DOCX extractor using python-docx, with graceful fallback."""

from __future__ import annotations

from pathlib import Path

from epe.ingestion.extractors import ExtractedDocument


def extract(path: Path) -> ExtractedDocument:
    warnings: list[str] = []
    try:
        import docx  # type: ignore[import-not-found]
    except ImportError as e:  # pragma: no cover
        warnings.append(f"python-docx not installed: {e}. Falling back to raw-text read.")
        return ExtractedDocument(
            path=path,
            text=path.read_text(encoding="utf-8", errors="replace"),
            extractor="docx-fallback",
            warnings=warnings,
        )

    document = docx.Document(path)
    parts: list[str] = []
    tables: list[list[list[str]]] = []
    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in document.tables:
        rows: list[list[str]] = []
        for row in table.rows:
            rows.append([cell.text for cell in row.cells])
        tables.append(rows)
        parts.append("\n".join(" | ".join(r) for r in rows))

    return ExtractedDocument(
        path=path,
        text="\n\n".join(parts),
        structure={"paragraph_count": len(document.paragraphs)},
        sections=[{"level": 0, "heading": "", "lines": parts}],
        tables=tables,
        extractor="docx",
        warnings=warnings,
    )
